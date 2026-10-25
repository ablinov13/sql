# Генерирует отчёты reports/Лаба*.docx на основе образца "Лаба2 (2).docx":
# титульный лист, оглавление и стили берутся из образца, содержимое — из этого файла.
# Запуск из корня решения: python tools/make_reports.py
import copy
import re
from pathlib import Path

import docx
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "Лаба2 (2).docx"
OUT = ROOT / "reports"
SHOTS = ROOT / "screenshots"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def q(tag):
    return "{%s}%s" % (W, tag)


class Report:
    """Собирает тело отчёта из абзацев-образцов шаблона."""

    def __init__(self, lab_number):
        self.doc = docx.Document(str(TEMPLATE))
        body = self.doc.element.body
        kids = list(body)
        # 0..28 — титульный лист, 29 — оглавление, 30 — разрыв страницы; дальше — содержимое образца
        self.proto = {
            "h_first": copy.deepcopy(kids[31]),
            "h": copy.deepcopy(kids[39]),
            "p": copy.deepcopy(kids[33]),
            "p_item": copy.deepcopy(kids[42]),
            "num": copy.deepcopy(kids[34]),
            "pic": copy.deepcopy(kids[54]),
            "caption": copy.deepcopy(kids[53]),
            "break": copy.deepcopy(kids[57]),
        }
        for el in kids[31:-1]:  # всё содержимое образца, кроме sectPr
            body.remove(el)
        self.body = body
        self.sect = kids[-1]
        self.bookmark = 100
        self.figure = 0
        self.first_heading = True
        self._set_lab_number(kids[9], lab_number)

    # --- титульный лист ---
    @staticmethod
    def _set_lab_number(par, number):
        runs = par.findall("w:r", NS)
        runs[-1].find("w:t", NS).text = str(number)

    # --- базовые операции ---
    def _append(self, el):
        self.sect.addprevious(el)
        return el

    @staticmethod
    def _set_text(p, text):
        """Оставляет в абзаце один run (с форматированием первого) и пишет в него текст."""
        runs = p.findall("w:r", NS)
        keep = runs[0]
        for r in runs[1:]:
            p.remove(r)
        for child in list(keep):
            if child.tag != q("rPr"):
                keep.remove(child)
        t = etree.SubElement(keep, q("t"))
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for tag in ("bookmarkStart", "bookmarkEnd"):
            for b in p.findall("w:" + tag, NS):
                p.remove(b)
        return keep

    def heading(self, text):
        if not self.first_heading:
            self.page_break()
        p = copy.deepcopy(self.proto["h_first"] if self.first_heading else self.proto["h"])
        self.first_heading = False
        run = self._set_text(p, text)
        self.bookmark += 1
        bs = etree.Element(q("bookmarkStart"))
        bs.set(q("id"), str(self.bookmark))
        bs.set(q("name"), "_Toc9%07d" % self.bookmark)
        be = etree.Element(q("bookmarkEnd"))
        be.set(q("id"), str(self.bookmark))
        run.addprevious(bs)
        run.addnext(be)
        self._append(p)

    def par(self, text):
        p = copy.deepcopy(self.proto["p"])
        self._set_text(p, text)
        self._append(p)

    def tasks(self, items):
        """Нумерованный список задач (1. 2. 3.), как в разделе «Цель и задачи» образца."""
        for it in items:
            p = copy.deepcopy(self.proto["num"])
            for r in p.findall("w:r", NS)[1:]:
                p.remove(r)
            self._set_text(p, it)
            self._append(p)

    def items(self, items):
        """Перечисление вида «model – модель автомобиля;», как в разделе «Ход работы» образца."""
        for it in items:
            p = copy.deepcopy(self.proto["p_item"])
            self._set_text(p, it)
            self._append(p)

    def bullets(self, items):
        """Маркированный список (как в выводе образца)."""
        for it in items:
            p = copy.deepcopy(self.proto["num"])
            for r in p.findall("w:r", NS)[1:]:
                p.remove(r)
            p.find("w:pPr/w:numPr/w:numId", NS).set(q("val"), "8")
            self._set_text(p, it)
            self._append(p)

    def picture(self, name, caption, max_w=16.0, max_h=21.0, crop_width=None):
        from PIL import Image
        path = SHOTS / name
        if crop_width:  # оставить только левую часть окна (например, редактор SQL и результат в DB Browser)
            img = Image.open(path)
            path = ROOT / "obj" / ("crop_" + name)
            img.crop((0, 0, crop_width, img.size[1])).save(path)
        w, h = Image.open(path).size
        width = min(max_w, max_h * w / h)
        p = self._append(copy.deepcopy(self.proto["pic"]))
        Paragraph(p, self.doc._body).add_run().add_picture(str(path), width=Cm(width))
        self.figure += 1
        c = copy.deepcopy(self.proto["caption"])
        self._set_text(c, f"Рисунок {self.figure} – {caption}")
        self._append(c)

    def page_break(self):
        self._append(copy.deepcopy(self.proto["break"]))

    def code(self, title, text):
        """Листинг: Courier New 10 пт, одинарный интервал (чтобы код не растягивался на десятки страниц)."""
        if title:
            p = copy.deepcopy(self.proto["p"])
            r = self._set_text(p, title)
            etree.SubElement(r.find("w:rPr", NS), q("b"))
            self._append(p)
        for line in text.rstrip("\n").split("\n"):
            p = etree.Element(q("p"))
            ppr = etree.SubElement(p, q("pPr"))
            sp = etree.SubElement(ppr, q("spacing"))
            sp.set(q("after"), "0")
            sp.set(q("line"), "240")
            sp.set(q("lineRule"), "auto")
            r = etree.SubElement(p, q("r"))
            rpr = etree.SubElement(r, q("rPr"))
            f = etree.SubElement(rpr, q("rFonts"))
            for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                f.set(q(a), "Courier New")
            sz = etree.SubElement(rpr, q("sz"))
            sz.set(q("val"), "20")
            t = etree.SubElement(r, q("t"))
            t.text = line.replace("\t", "    ")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            self._append(p)

    def save(self, filename):
        OUT.mkdir(exist_ok=True)
        path = OUT / filename
        self.doc.save(str(path))
        return path


def src(rel):
    return (ROOT / rel).read_text(encoding="utf-8")


def appendix(r, files):
    r.heading("Приложение А")
    for i, f in enumerate(files, 1):
        r.code(f"Листинг А.{i} – {f}" if len(files) > 1 else "", src(f))
        if i < len(files):
            r.par("")


TASK_LAB1 = ["Скачать любой вариант клиента SQLite;",
             "Создать таблицу и загрузить данные из предыдущего задания, используя язык SQL;",
             "Выполнить 2–3 запроса SELECT."]

# ---------------------------------------------------------------- Лаба 1
def lab1():
    r = Report(1)
    r.heading("Цель и задачи")
    r.par("Целью данной лабораторной работы является выполнение запросов поиска данных с помощью языка SQL.")
    r.par("Для достижения цели были поставлены следующие задачи:")
    r.tasks(TASK_LAB1)
    r.heading("Ход работы")
    r.par("В рамках предыдущего задания была выбрана предметная область «Книжный магазин» и подготовлен "
          "текстовый файл books.txt, содержащий 25 записей о книгах. Значения в каждой строке разделены запятыми.")
    r.par("В качестве клиента SQLite был скачан и установлен DB Browser for SQLite – программа с графическим "
          "интерфейсом для работы с базами данных SQLite. В ней была создана база данных lab1.db.")
    r.par("Создание таблицы и загрузка данных выполнялись с помощью языка SQL: был написан скрипт lab1.sql "
          "(приложение А), содержащий оператор CREATE TABLE и оператор INSERT INTO с 25 наборами значений. "
          "Скрипт выполнен на вкладке «Execute SQL».")
    r.par("Структура таблицы books включает следующие поля:")
    r.items(["title – название книги (TEXT);",
             "genre – жанр (TEXT): Роман, Классика, Фантастика, Детектив, Антиутопия, Фэнтези, Научпоп;",
             "price – цена в рублях (INTEGER);",
             "year – год издания (INTEGER);",
             "author – автор (TEXT)."])
    r.par("Всего в таблицу было загружено 25 записей.")
    r.par("Следующим этапом было выполнение запросов. Всего их было три.")
    r.par("Первый запрос: SELECT title, price FROM books WHERE genre = 'Детектив' AND price < 450. "
          "Нахождение детективов стоимостью менее 450 рублей. Результат запроса представлен на рисунке 1.")
    r.picture("lab1_q1.png", "Первый запрос", crop_width=648)
    r.par("Второй запрос: SELECT title, genre, price FROM books ORDER BY price DESC LIMIT 5. "
          "Вывод пяти самых дорогих книг. Результат запроса представлен на рисунке 2.")
    r.picture("lab1_q2.png", "Второй запрос", crop_width=648)
    r.par("Третий запрос: SELECT genre, COUNT(*) AS books_count, CAST(ROUND(AVG(price)) AS INTEGER) AS avg_price "
          "FROM books GROUP BY genre ORDER BY books_count DESC, genre. Подсчёт количества книг и средней цены "
          "по каждому жанру. Результат запроса представлен на рисунке 3.")
    r.picture("lab1_q3.png", "Третий запрос", crop_width=648)
    r.heading("Вывод")
    r.par("В ходе лабораторной работы были выполнены все задачи: скачан клиент SQLite (DB Browser for SQLite), "
          "с помощью языка SQL создана таблица books и загружены в неё 25 записей о книгах, а также выполнены "
          "три запроса SELECT.")
    r.par("С помощью запросов удалось:")
    r.bullets(["найти детективы стоимостью менее 450 рублей;",
               "вывести топ-5 самых дорогих книг;",
               "определить количество книг и среднюю цену в каждом жанре."])
    r.par("В итоге были освоены основные операции работы с базой данных SQLite: создание таблицы, загрузка "
          "данных оператором INSERT, выполнение запросов с условиями WHERE, сортировкой ORDER BY, "
          "ограничением LIMIT, группировкой GROUP BY и агрегатными функциями COUNT и AVG.")
    appendix(r, ["lab1/lab1.sql"])
    return r.save("Лаба1.docx")


# ---------------------------------------------------------------- Лаба 2
def lab2():
    r = Report(2)
    r.heading("Цель и задачи")
    r.par("Целью данной лабораторной работы является подключение к встраиваемой базе данных из программы, "
          "создание таблицы, загрузка данных и выполнение запросов поиска данных с помощью языка SQL.")
    r.par("Для достижения цели были поставлены следующие задачи:")
    r.tasks(["Скачать драйвер для встраиваемой базы данных;",
             "Подключить драйвер к проекту (среда разработки: Eclipse, NetBeans, IDEA или другая);",
             "Подключиться к базе из программы;",
             "Создать таблицу и загрузить данные из предыдущих заданий (с помощью SQL);",
             "Выполнить 2–3 запроса SELECT."])
    r.heading("Ход работы")
    r.par("Для выполнения лабораторной работы был использован язык программирования C++ и среда разработки "
          "Visual Studio 2022. В качестве встраиваемой базы данных применялась SQLite – она не требует "
          "установки отдельного сервера, а вся база хранится в одном файле.")
    r.par("Роль драйвера (аналога JDBC-драйвера H2 для Java) выполняет библиотека SQLite. С официального сайта "
          "sqlite.org была скачана amalgamation-версия библиотеки (файлы sqlite3.c и sqlite3.h). Для неё в решении "
          "создан отдельный проект sqlite, собирающий статическую библиотеку sqlite.lib, а проект lab2 подключает "
          "её через ссылку на проект и путь к заголовочному файлу.")
    r.par("Подключение к базе выполняется функцией sqlite3_open, которая открывает файл lab2.db или создаёт его. "
          "В базе создана таблица books со следующими полями:")
    r.items(["title – название книги (TEXT);",
             "genre – жанр (TEXT);",
             "price – цена в рублях (INTEGER);",
             "year – год издания (INTEGER);",
             "author – автор (TEXT)."])
    r.par("Данные были загружены из текстового файла books.txt, содержащего 25 записей о книгах. Загрузка "
          "выполнялась подготовленным SQL-запросом INSERT INTO books VALUES (?, ?, ?, ?, ?): для каждой строки "
          "файла значения подставлялись функциями sqlite3_bind_text и sqlite3_bind_int, а запрос выполнялся "
          "функцией sqlite3_step. Все вставки выполнены в одной транзакции (BEGIN … COMMIT).")
    r.par("После загрузки данных было выполнено три запроса SELECT.")
    r.par("Первый запрос: SELECT title, price FROM books WHERE genre = 'Детектив' AND price < 450. "
          "Нахождение детективов стоимостью менее 450 рублей.")
    r.par("Второй запрос: SELECT title, genre, price FROM books ORDER BY price DESC LIMIT 5. "
          "Вывод пяти самых дорогих книг.")
    r.par("Третий запрос: SELECT title, author, year FROM books WHERE year = 2022 ORDER BY title. "
          "Нахождение книг, изданных в 2022 году.")
    r.par("Результаты выполнения всех трёх запросов представлены на рисунке 1.")
    r.picture("lab2.png", "Результаты выполнения запросов")
    r.par("Все запросы выполнены без ошибок. Программа успешно подключается к базе данных, создаёт таблицу, "
          "загружает данные и выводит результаты поиска в консоль.")
    r.heading("Вывод")
    r.par("В ходе лабораторной работы были выполнены все задачи: подключена библиотека встраиваемой базы данных "
          "SQLite к проекту на C++ в Visual Studio 2022, выполнено подключение к базе из программы, создана "
          "таблица books и загружены в неё 25 записей о книгах, а также выполнены три запроса SELECT.")
    r.par("С помощью запросов удалось:")
    r.bullets(["найти детективы стоимостью менее 450 рублей;",
               "вывести топ-5 самых дорогих книг;",
               "найти все книги 2022 года издания."])
    r.par("В итоге были освоены основные операции работы с базой данных SQLite из программы на C++: подключение "
          "к БД, создание таблицы, загрузка данных из файла с помощью подготовленных запросов, выполнение "
          "запросов с условиями WHERE, сортировкой ORDER BY, ограничением LIMIT и чтение результатов. "
          "Полученные навыки позволяют интегрировать работу с базами данных в прикладные программы.")
    appendix(r, ["lab2/main.cpp"])
    return r.save("Лаба2.docx")


# ---------------------------------------------------------------- Лаба 3.1
def lab3():
    r = Report("3.1")
    r.heading("Цель и задачи")
    r.par("Целью данной лабораторной работы является перевод программы работы с базой данных из лабораторной "
          "работы 2 на шаблон проектирования DAO (Data Access Object).")
    r.par("Для достижения цели были поставлены следующие задачи:")
    r.tasks(["Перевести код из работы 2 в код с API DAO;",
             "Добавление элементов в базу: создавать объект и помещать его в базу через XxxDao.insert;",
             "Выборка элементов: через XxxDao.getXxx, итерация по списку объектов;",
             "Не использовать вывод на экран внутри методов XxxDao, не перехватывать в них исключения, а "
             "передавать их наверх и обрабатывать в функции main."])
    r.heading("Ход работы")
    r.par("Работа выполнена на языке C++ в среде Visual Studio 2022 с базой данных SQLite. Для работы DAO "
          "с двумя таблицами структура базы была расширена: кроме таблицы книг добавлена таблица авторов, "
          "а каждая книга ссылается на своего автора по внешнему ключу.")
    r.par("Структура таблицы authors:")
    r.items(["id – идентификатор автора (INTEGER, первичный ключ);",
             "name – имя автора (TEXT, уникальное);",
             "country – страна (TEXT)."])
    r.par("Структура таблицы books:")
    r.items(["id – идентификатор книги (INTEGER, первичный ключ);",
             "title – название (TEXT);",
             "genre – жанр (TEXT);",
             "price – цена в рублях (INTEGER);",
             "year – год издания (INTEGER);",
             "author_id – идентификатор автора (INTEGER, внешний ключ на authors.id)."])
    r.par("Программа разделена на следующие части:")
    r.items(["Models.h – структуры Author и Book, соответствующие строкам таблиц;",
             "Database – класс подключения к базе (открывает соединение в конструкторе и закрывает в "
             "деструкторе), класс Statement – обёртка подготовленного запроса, класс исключения DbException;",
             "AuthorDao – методы dropTable, createTable, insert, getAll, getByName для таблицы authors;",
             "BookDao – методы dropTable, createTable, insert, getAll, getByGenre, getByPriceRange, "
             "getAllWithAuthors для таблицы books;",
             "main.cpp – чтение исходных файлов, вызовы методов DAO, вывод результатов и обработка ошибок."])
    r.par("Добавление элементов выполняется так: для каждой строки файла создаётся объект Author или Book, "
          "заполняются его поля, и объект передаётся в метод authorDao.insert или bookDao.insert. Чтобы связать "
          "книгу с автором, идентификатор автора находится методом authorDao.getByName.")
    r.par("Выборка элементов выполняется методами getXxx, которые возвращают список объектов (std::vector); "
          "в функции main выполняется итерация по этому списку и вывод каждого объекта.")
    r.par("Согласно замечаниям, внутри методов DAO нет вывода на экран и нет перехвата исключений: при ошибке "
          "SQLite методы выбрасывают исключение DbException, которое передаётся наверх и обрабатывается в функции "
          "main. Весь вывод результатов также выполняется только в main.")
    r.par("Были выполнены следующие запросы:")
    r.items(["bookDao.getByGenre(\"Фантастика\") – книги жанра «Фантастика»;",
             "bookDao.getByPriceRange(500, 700) – книги стоимостью от 500 до 700 рублей;",
             "bookDao.getAllWithAuthors() – книги вместе с именами авторов (соединение двух таблиц JOIN)."])
    r.par("Результат работы программы представлен на рисунке 1.")
    r.picture("lab3.png", "Результат работы программы с DAO")
    r.heading("Вывод")
    r.par("В ходе лабораторной работы код из лабораторной работы 2 был переведён на API DAO. Создана база данных "
          "из двух связанных таблиц authors и books, для каждой из них реализован свой класс DAO.")
    r.par("В результате:")
    r.bullets(["добавление данных выполняется через создание объектов и вызов XxxDao.insert;",
               "выборка выполняется через методы XxxDao.getXxx с итерацией по списку объектов;",
               "методы DAO не выводят сообщений и не перехватывают исключения – вывод и обработка ошибок "
               "сосредоточены в функции main."])
    r.par("Использование DAO отделяет код работы с базой данных от логики программы: весь SQL собран в классах "
          "DAO, а основная программа работает с обычными объектами.")
    appendix(r, ["lab3/Models.h", "lab3/Database.h", "lab3/Database.cpp", "lab3/AuthorDao.h",
                 "lab3/AuthorDao.cpp", "lab3/BookDao.h", "lab3/BookDao.cpp", "lab3/main.cpp"])
    return r.save("Лаба3.docx")


# ---------------------------------------------------------------- Лаба 4
def lab4():
    r = Report(4)
    r.heading("Цель и задачи")
    r.par("Целью данной лабораторной работы является перевод программы работы с базой данных на механизм ORM "
          "(объектно-реляционного отображения).")
    r.par("Для достижения цели были поставлены следующие задачи:")
    r.tasks(["Перевести код лабораторной работы 3 (DAO с двумя таблицами) на механизм ORM;",
             "Использовать любой ORM-фреймворк;",
             "Создать структуру базы;",
             "Добавить элементы в таблицы (20–30 элементов, колонки – строки и числа, с возможностью "
             "группировать по повторяющимся значениям и диапазонам);",
             "Выполнить 2–3 осмысленных запроса."])
    r.heading("Ход работы")
    r.par("Работа выполнена на языке C++ в среде Visual Studio 2022. Рекомендованный фреймворк ORMLite "
          "предназначен только для Java, поэтому в качестве ORM для C++ была выбрана библиотека sqlite_orm. "
          "Она состоит из одного заголовочного файла sqlite_orm.h и работает поверх SQLite.")
    r.par("Используются те же структуры Author и Book, что и в лабораторной работе 3. Структура базы описывается "
          "в функции makeStorage (файл Storage.h): функция make_table задаёт таблицу, make_column – соответствие "
          "колонки таблицы полю структуры, primary_key().autoincrement() – первичный ключ, "
          "foreign_key(&Book::authorId).references(&Author::id) – внешний ключ на таблицу авторов. Таблицы "
          "создаются вызовом storage.sync_schema(), SQL-команды CREATE TABLE библиотека формирует сама.")
    r.par("Добавление элементов выполняется методом storage.insert(объект) внутри транзакции "
          "storage.transaction. В таблицу authors добавлено 11 авторов, в таблицу books – 25 книг. "
          "Колонки содержат строки (название, жанр, имя автора, страна) и числа (цена, год), жанры и страны "
          "повторяются, что позволяет группировать данные.")
    r.par("Были выполнены три осмысленных запроса средствами ORM:")
    r.items(["группировка по повторяющимся значениям – количество книг и средняя цена в каждом жанре: "
             "select(columns(&Book::genre, count(&Book::id), avg(&Book::price)), group_by(&Book::genre));",
             "выборка по диапазону – книги стоимостью от 500 до 700 рублей: "
             "get_all<Book>(where(between(&Book::price, 500, 700)), order_by(&Book::price));",
             "запрос к двум таблицам – книги авторов из Великобритании: select(columns(&Author::name, "
             "&Book::title, &Book::year), inner_join<Author>(on(c(&Book::authorId) == &Author::id)), "
             "where(c(&Author::country) == \"Великобритания\"))."])
    r.par("Результат работы программы представлен на рисунке 1.")
    r.picture("lab4.png", "Результат работы программы с ORM")
    r.par("Как и в лабораторной работе 3, вывод результатов и обработка ошибок выполняются только в функции "
          "main; ошибки базы данных библиотека sqlite_orm сообщает исключением std::system_error.")
    r.heading("Вывод")
    r.par("В ходе лабораторной работы код лабораторной работы 3 был переведён на механизм ORM с использованием "
          "библиотеки sqlite_orm. Создана структура базы из двух таблиц, добавлены 11 авторов и 25 книг, "
          "выполнены три осмысленных запроса.")
    r.par("С помощью запросов удалось:")
    r.bullets(["определить количество книг и среднюю цену в каждом жанре;",
               "найти книги в диапазоне цен от 500 до 700 рублей;",
               "вывести книги британских авторов, соединив таблицы books и authors."])
    r.par("По сравнению с DAO механизм ORM избавляет от написания SQL-запросов и ручного переноса значений "
          "колонок в поля объектов: достаточно один раз описать соответствие структур и таблиц, после чего "
          "библиотека сама формирует запросы и возвращает готовые объекты.")
    appendix(r, ["lab4/Models.h", "lab4/Storage.h", "lab4/main.cpp"])
    return r.save("Лаба4.docx")


# ---------------------------------------------------------------- Лаба 5
def lab5():
    r = Report("5")
    r.heading("Цель и задачи")
    r.par("Целью данной лабораторной работы является изучение работы HTTP-сервера и формирование HTML-страницы "
          "с данными из предыдущих лабораторных работ.")
    r.par("Для достижения цели были поставлены следующие задачи:")
    r.tasks(["Запустить демо http-hello.js;",
             "Запустить демо http-hello-html.js;",
             "Сверстать статическую страницу – таблицу с данными из предыдущих лабораторных работ, вернуть её "
             "в браузер по схеме http-hello-html."])
    r.heading("Ход работы")
    r.par("Демонстрационные программы http-hello.js и http-hello-html.js запускались на платформе Node.js "
          "командами node http-hello.js и node http-hello-html.js. Обе программы создают HTTP-сервер на порту "
          "3000: первая возвращает текст «Hello World», вторая – HTML-страницу с результатом выражения 3 + 5. "
          "Результаты в браузере по адресу http://localhost:3000/ представлены на рисунках 1 и 2.")
    r.picture("lab5_js_hello.png", "Результат работы http-hello.js")
    r.picture("lab5_js_html.png", "Результат работы http-hello-html.js")
    r.par("Для лучшего понимания схемы работы сервера обе демонстрационные программы были повторены на языке "
          "C++ (проекты lab5_1 и lab5_2 в Visual Studio 2022). Сервер написан на сокетах Windows (WinSock): "
          "функции socket, bind и listen соответствуют вызову http.createServer(...).listen(3000), цикл accept – "
          "функции обработки запроса, а функция send отправляет браузеру ответ: строку статуса "
          "HTTP/1.1 200 OK, заголовки Content-Type и Content-Length и тело ответа. Результаты представлены "
          "на рисунках 3 и 4.")
    r.picture("lab5_1.png", "Результат работы lab5_1 (аналог http-hello.js на C++)")
    r.picture("lab5_2.png", "Результат работы lab5_2 (аналог http-hello-html.js на C++)")
    r.par("Статическая страница с таблицей книг (проект lab5_3) сформирована по схеме http-hello-html: "
          "сначала в ответ записывается начало HTML-документа (DOCTYPE, head с кодировкой utf-8 и стилями "
          "таблицы, начало body и заголовок таблицы), затем строки таблицы с данными, затем закрывающие теги. "
          "Данные – 25 книг из файла books.txt, использованного в предыдущих лабораторных работах. Таблица "
          "содержит колонки: номер, название, автор, жанр, цена и год. Специальные символы HTML в данных "
          "экранируются функцией escapeHtml.")
    r.par("Окно запущенного сервера показано на рисунке 5, страница в браузере – на рисунке 6.")
    r.picture("lab5_3_console.png", "Запущенный сервер lab5_3")
    r.picture("lab5_3.png", "Таблица книг в браузере")
    r.heading("Вывод")
    r.par("В ходе лабораторной работы были запущены демонстрационные программы http-hello.js и "
          "http-hello-html.js, их работа повторена на языке C++, а также сверстана статическая страница "
          "с таблицей данных из предыдущих лабораторных работ, которая возвращается в браузер по схеме "
          "http-hello-html.")
    r.par("В результате были изучены:")
    r.bullets(["структура HTTP-запроса и HTTP-ответа (строка статуса, заголовки, тело);",
               "назначение заголовка Content-Type и указания кодировки utf-8 для отображения русского текста;",
               "формирование HTML-страницы на сервере по частям: начало документа, данные, окончание."])
    appendix(r, ["lab5/http-hello.js", "lab5/http-hello-html.js", "lab5_1/main.cpp", "lab5_2/main.cpp",
                 "lab5_3/main.cpp"])
    return r.save("Лаба5.docx")


if __name__ == "__main__":
    for f in (lab1, lab2, lab3, lab4, lab5):
        print("saved", f().name)
